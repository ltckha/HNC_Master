"""Đối chiếu cờ shopee_upload / tiktok_upload theo file Liên kết (Cập nhật hàng loạt) tải từ Seller Center — anh Kha "làm" 29/10.
Cùng luật web public/tools/lien-ket-san.js (build: cột theo tên hàng 1, ghép SKU → mã SP) và bot dat_lai_co_san:
  - set   = mã có trên sàn mà ô cờ chưa "Done…" (trống / cờ lỗi) → "Done - <giờ>"; ô đã Done giữ nguyên
  - unset = ô cờ có chữ mà mã KHÔNG có trên sàn → xoá ô
Từ 30/10 (anh Kha): menu Mac Đăng mẫu mới Shopee / TikTok coi "chưa đăng" = mã không có trong file sàn mới nhất (file_moi_nhat + ma_tren_san), không đọc / ghi cờ.
Dùng từ main.py (Đăng mẫu mới; menu [9] đối chiếu cờ cũ); chạy riêng chỉ đọc:
  venv/bin/python tools/lien_ket_san.py shopee|tiktok <file.xlsx|.csv>
"""
import csv
import os
import re
import sys

REQ = {'shopee': ['et_title_product_id', 'et_title_variation_sku'], 'tiktok': ['product_id', 'seller_sku']}
SIZE_DUOI = re.compile(r'^\d{2}(\.5)?$')


def doc_file(path, kenh):
    """→ list hàng (list ô chuỗi). Shopee: sheet đầu; TikTok: sheet 'Template' (không có thì sheet đầu)."""
    if path.lower().endswith('.csv'):
        with open(path, encoding='utf-8-sig', newline='') as f:
            return [list(r) for r in csv.reader(f)]
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Template'] if kenh == 'tiktok' and 'Template' in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = [['' if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def ma_tren_san(kenh, rows, codes):
    """Như web build(): → (tập mã có trên sàn, số dòng có mã sàn, các SKU không khớp)."""
    head = [str(h or '').strip() for h in (rows[0] if rows else [])]
    thieu = [c for c in REQ[kenh] if c not in head]
    if thieu:
        raise ValueError(f"File thiếu cột {', '.join(thieu)} (sai sàn hoặc mẫu file đã đổi)")
    ix = {n: i for i, n in enumerate(head)}
    tap = set(codes)
    theo_dai = sorted(codes, key=len, reverse=True)

    def g(r, *ns):
        for n in ns:
            i = ix.get(n)
            if i is not None and i < len(r) and str(r[i] or '').strip():
                return str(r[i]).strip()
        return ''

    def to_code(sku, vn):
        sku = sku.upper()
        if not sku:
            return None
        if sku in tap:
            return sku
        tail = vn.split(',')[-1].strip().upper() if ',' in vn else ''
        if tail and sku.endswith(tail) and sku[:-len(tail)] in tap:
            return sku[:-len(tail)]
        for c in theo_dai:
            if sku.startswith(c) and sku != c and SIZE_DUOI.match(sku[len(c):]):
                return c
        return None

    co, n, sai = set(), 0, []
    for r in rows[1:]:
        if not r:
            continue
        item = g(r, 'et_title_product_id') if kenh == 'shopee' else g(r, 'product_id')
        if not re.fullmatch(r'\d{5,25}', item):
            continue  # hàng hướng dẫn / tiêu đề tiếng Việt
        n += 1
        sku = g(r, 'et_title_variation_sku', 'et_title_parent_sku') if kenh == 'shopee' else g(r, 'seller_sku')
        code = to_code(sku, g(r, 'et_title_variation_name') if kenh == 'shopee' else g(r, 'variation_value'))
        if code:
            co.add(code)
        else:
            sai.append(sku or '(trống)')
    return co, n, sorted(set(sai))


def file_moi_nhat(kenh, thu_muc=os.path.expanduser('~/Downloads'), xem=40):
    """File Liên kết / Cập nhật hàng loạt mới nhất của sàn trong thư mục (theo giờ sửa, đủ cột REQ) → đường dẫn hoặc None.
    Anh Kha 30/10: "chưa đăng sàn" = mã KHÔNG có trong file sàn mới nhất anh tải về (không dùng cờ shopee_upload / tiktok_upload)."""
    try:
        ds = [os.path.join(thu_muc, f) for f in os.listdir(thu_muc)
              if f.lower().endswith(('.xlsx', '.csv')) and not f.startswith(('~$', '.'))]
    except OSError:
        return None
    for p in sorted(ds, key=os.path.getmtime, reverse=True)[:xem]:
        try:
            rows = doc_file(p, kenh)
        except Exception:
            continue
        head = [str(h or '').strip() for h in (rows[0] if rows else [])]
        if all(c in head for c in REQ[kenh]):
            return p
    return None


def doi_chieu(co_tren_san, co_sheet):
    """co_sheet = {mã: giá trị ô cờ} → (set, unset) như bot dat_lai_co_san."""
    dat = sorted(c for c, v in co_sheet.items() if c in co_tren_san and not str(v).strip().lower().startswith('done'))
    go = sorted(c for c, v in co_sheet.items() if c not in co_tren_san and str(v).strip())
    return dat, go


if __name__ == '__main__':
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from tools.tach_cot_sanpham import _mo
    from tools.tu_dien_sanpham import SHEET_ID
    kenh, path = sys.argv[1], sys.argv[2]
    sp = _mo(SHEET_ID).get_all_values()
    h = [x.strip() for x in sp[1]]
    ic, iv = h.index('product_code'), h.index(f'{kenh}_upload')
    co_sheet = {r[ic].strip().upper(): (r[iv] if iv < len(r) else '') for r in sp[2:] if r[ic].strip()}
    co, n, sai = ma_tren_san(kenh, doc_file(path, kenh), list(co_sheet))
    dat, go = doi_chieu(co, co_sheet)
    print(f"{kenh}: {n} dòng có mã sàn → {len(co)} mã khớp · {len(sai)} SKU không khớp · Sheet đang ghi {sum(1 for v in co_sheet.values() if str(v).strip())} mã")
    print(f"  đặt Done {len(dat)}: {', '.join(dat[:10])}{' …' if len(dat) > 10 else ''}")
    print(f"  gỡ cờ {len(go)}: {', '.join(go[:10])}{' …' if len(go) > 10 else ''}")
